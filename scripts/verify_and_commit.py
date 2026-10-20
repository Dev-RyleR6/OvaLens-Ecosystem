#!/usr/bin/env python
"""
verify_and_commit.py — OvaLens Automated Quality Gate, Task Checklist Updater, & Git Commit Utility

Usage:
    python scripts/verify_and_commit.py --task TASK-BE-AUTH-01 -m "feat(backend): implement auth login endpoint"
    python scripts/verify_and_commit.py --task TASK-EDGE-TRK-01 -m "feat(edge): add cradle tracking queue" [--skip-frontend]
"""

import sys
from io import TextIOWrapper

if isinstance(sys.stdout, TextIOWrapper):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import os
import re
import subprocess
import argparse
from datetime import datetime

VALID_TYPES = ["feat", "fix", "refactor", "perf", "docs", "test", "chore"]
VALID_SCOPES = ["backend", "edge", "dashboard", "firmware", "docs", "agents", "ci", "repo", "architecture"]

def run_command(cmd, cwd=None, exit_on_fail=True):
    print(f"\n[RUNNING] {cmd} (cwd: {cwd or '.'})")
    res = subprocess.run(cmd, shell=True, cwd=cwd)
    if res.returncode != 0 and exit_on_fail:
        print(f"\n❌ FAILED: {cmd} exited with return code {res.returncode}")
        sys.exit(res.returncode)
    return res.returncode

def parse_commit_message(args):
    if args.message:
        pattern = r"^(feat|fix|refactor|perf|docs|test|chore)\(([a-z0-9_-]+)\):\s*(.+)$"
        m = re.match(pattern, args.message.strip())
        if not m:
            print("\n❌ Invalid commit message format!")
            print("Must follow Conventional Commits: <type>(<scope>): <subject>")
            print(f"Example: feat(backend): add penoy yield calculation")
            sys.exit(1)
        commit_type, commit_scope, commit_subject = m.groups()
    else:
        if not args.type or not args.scope or not args.msg:
            print("\n❌ Either --message (-m) OR (--type, --scope, --msg) must be provided!")
            sys.exit(1)
        commit_type = args.type.lower()
        commit_scope = args.scope.lower()
        commit_subject = args.msg.strip()

    if commit_type not in VALID_TYPES:
        print(f"\n❌ Invalid type '{commit_type}'. Must be one of: {', '.join(VALID_TYPES)}")
        sys.exit(1)

    full_message = f"{commit_type}({commit_scope}): {commit_subject}"
    return commit_type, commit_scope, commit_subject, full_message

def update_changelog(repo_root, commit_type, commit_scope, commit_subject, task_id=None):
    changelog_path = os.path.join(repo_root, "CHANGELOG.md")
    if not os.path.exists(changelog_path):
        return

    type_mapping = {
        "feat": "Added",
        "fix": "Fixed",
        "perf": "Changed",
        "refactor": "Changed",
        "docs": "Added",
        "chore": "Changed",
        "test": "Added"
    }
    section_name = type_mapping.get(commit_type, "Changed")
    task_tag = f"[{task_id.upper()}] " if task_id else ""
    entry = f"- **{commit_scope.capitalize()}**: {task_tag}{commit_subject} ({datetime.now().strftime('%Y-%m-%d')})"

    with open(changelog_path, "r", encoding="utf-8") as f:
        content = f.read()

    target_heading = f"### {section_name}"
    if target_heading in content:
        replacement = f"{target_heading}\n{entry}"
        content = content.replace(target_heading, replacement, 1)
        with open(changelog_path, "w", encoding="utf-8") as f:
            f.write(content)
        print(f"📝 Updated CHANGELOG.md under [{section_name}]")

def update_checklists(repo_root, task_id):
    if not task_id:
        return
    task_clean = task_id.strip().upper()

    # 1. Update PROJECT_PLAN.md
    plan_path = os.path.join(repo_root, "PROJECT_PLAN.md")
    if os.path.exists(plan_path):
        with open(plan_path, "r", encoding="utf-8") as f:
            lines = f.readlines()
        updated = False
        for i, line in enumerate(lines):
            if task_clean.lower() in line.lower() and "- [ ]" in line:
                lines[i] = line.replace("- [ ]", "- [x]", 1)
                updated = True
                print(f"✅ Marked task as [x] in PROJECT_PLAN.md: {task_clean}")
                break
        if updated:
            with open(plan_path, "w", encoding="utf-8") as f:
                f.writelines(lines)

    # 2. Update Excel Workbook docs/OvaLens_Capstone_Feature_Checklist.xlsx
    excel_path = os.path.join(repo_root, "docs", "OvaLens_Capstone_Feature_Checklist.xlsx")
    if os.path.exists(excel_path):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            wb = openpyxl.load_workbook(excel_path)
            if "Master Checklist" in wb.sheetnames:
                ws = wb["Master Checklist"]
                status_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                status_font = Font(name="Segoe UI", size=10, bold=True, color="166534")
                matched = False
                for row in range(4, ws.max_row + 1):
                    val = str(ws.cell(row=row, column=1).value or "").strip().upper()
                    if task_clean == val or task_clean in val:
                        ws.cell(row=row, column=8, value="Done")
                        ws.cell(row=row, column=8).fill = status_fill
                        ws.cell(row=row, column=8).font = status_font
                        matched = True
                        print(f"✅ Marked task as 'Done' in Excel Checklist: {val}")
                        break
                if matched:
                    wb.save(excel_path)
        except Exception as e:
            print(f"⚠️ Warning: Could not update Excel file automatically: {e}")

def main():
    parser = argparse.ArgumentParser(description="OvaLens Automated Verification, Task Checklist, and Commit Tool")
    parser.add_argument("-m", "--message", help="Full commit message: <type>(<scope>): <subject>")
    parser.add_argument("-t", "--task", help="Task ID from PROJECT_PLAN.md to mark as completed (e.g. TASK-EDGE-TRK-01)")
    parser.add_argument("--type", choices=VALID_TYPES, help="Commit type")
    parser.add_argument("--scope", help="Subsystem scope")
    parser.add_argument("--msg", help="Commit subject line")
    parser.add_argument("--skip-tests", action="store_true", help="Skip running test suites (not recommended)")
    parser.add_argument("--skip-frontend", action="store_true", help="Skip frontend build check")
    parser.add_argument("--no-changelog", action="store_true", help="Skip updating CHANGELOG.md")
    parser.add_argument("--verify-only", action="store_true",
                        help="Run checks without modifying checklists, changelog, staging, or committing")
    parser.add_argument("--staged-only", action="store_true",
                        help="Commit the prepared index without staging unrelated workspace changes")
    args = parser.parse_args()

    repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    commit_type, commit_scope, commit_subject, full_message = parse_commit_message(args)

    print("=" * 70)
    print(f"🚀 OvaLens Automated Quality Gate & Commit Workflow")
    if args.task:
        print(f"Task ID: {args.task.upper()}")
    print(f"Commit:  {full_message}")
    print("=" * 70)

    # 1. Run Tests
    if not args.skip_tests:
        print("\n🧪 [1/3] Running Backend API Test Suite...")
        run_command(f"{sys.executable} -m pytest tests/test_api.py -v", cwd=os.path.join(repo_root, "backend"))

        print("\n🧪 [2/3] Running Edge CV Pipeline Test Suite...")
        run_command(f"{sys.executable} -m pytest tests -v", cwd=os.path.join(repo_root, "edge"))

        if not args.skip_frontend:
            print("\n🧪 [3/3] Checking React Dashboard Build...")
            run_command("npm.cmd run build", cwd=os.path.join(repo_root, "dashboard"))
    else:
        print("\n⚠️ Skipping automated test suites as requested.")

    if args.verify_only:
        print("\nSUCCESS: Verification complete; no files staged or committed.")
        return

    # 2. Update Checklist (PROJECT_PLAN.md & Excel)
    if args.task:
        print(f"\n📋 Updating Task Checklist for [{args.task.upper()}]...")
        update_checklists(repo_root, args.task)

    # 3. Update Changelog
    if not args.no_changelog:
        update_changelog(repo_root, commit_type, commit_scope, commit_subject, args.task)

    # 4. Git stage and commit
    print("\n📦 Staging files and creating Git commit...")
    if not args.staged_only:
        run_command("git add .", cwd=repo_root)
    run_command(f'git commit -m "{full_message}"', cwd=repo_root)

    print("\n" + "=" * 70)
    print(f"✅ SUCCESS: Verified, checked off, and committed: {full_message}")
    print("=" * 70)

if __name__ == "__main__":
    main()
